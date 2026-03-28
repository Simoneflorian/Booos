import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side,
)
from openpyxl.utils import get_column_letter


# Color palette
COLOR_HEADER_BG = "1E40AF"
COLOR_HEADER_FG = "FFFFFF"
COLOR_SUMMARY_BG = "EFF6FF"
COLOR_ROW_ALT = "F8FAFC"
COLOR_TOTAL_BG = "DBEAFE"
COLOR_BORDER = "CBD5E1"
COLOR_SEARCH_INPUT = "FEF08A"
COLOR_SEARCH_BORDER = "D97706"
COLOR_SEC_RECEIPT = "DBEAFE"
COLOR_SEC_ARTICLE = "DCFCE7"


def _border(style="thin"):
    side = Side(style=style, color=COLOR_BORDER)
    return Border(left=side, right=side, top=side, bottom=side)


def _amber_border():
    s = Side(style="medium", color=COLOR_SEARCH_BORDER)
    return Border(left=s, right=s, top=s, bottom=s)


def _header_font():
    return Font(name="Calibri", bold=True, color=COLOR_HEADER_FG, size=11)


def _bold(size=11):
    return Font(name="Calibri", bold=True, size=size)


def _normal(size=10):
    return Font(name="Calibri", size=size)


def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def _center():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)


def _left():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)


def _right():
    return Alignment(horizontal="right", vertical="center")


def _set_col_widths(ws, widths: dict):
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width


def create_excel(receipts: list) -> io.BytesIO:
    wb = Workbook()
    wb.remove(wb.active)

    # Sheet order: Suche first, then data sheets
    _create_search_sheet(wb, receipts)
    _create_summary_sheet(wb, receipts)
    _create_articles_sheet(wb, receipts)
    for idx, receipt in enumerate(receipts, start=1):
        _create_receipt_sheet(wb, receipt, idx)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ---------------------------------------------------------------------------
# Search Sheet  (Tab 1 — first thing the user sees)
# ---------------------------------------------------------------------------
def _create_search_sheet(wb: Workbook, receipts: list):
    ws = wb.create_sheet("Suche")
    ws.sheet_view.showGridLines = False
    ws.sheet_view.tabSelected = True  # activate this sheet on open

    # --- Title ---
    ws.row_dimensions[1].height = 38
    ws.merge_cells("A1:N1")
    t = ws["A1"]
    t.value = "Quittungssuche"
    t.font = Font(name="Calibri", bold=True, size=18, color=COLOR_HEADER_FG)
    t.fill = _fill(COLOR_HEADER_BG)
    t.alignment = _center()

    # --- Search input row ---
    ws.row_dimensions[3].height = 32
    lbl = ws.cell(row=3, column=1, value="Suchbegriff:")
    lbl.font = _bold(13)
    lbl.alignment = Alignment(horizontal="right", vertical="center")

    inp = ws.cell(row=3, column=2, value="")
    inp.fill = _fill(COLOR_SEARCH_INPUT)
    inp.font = Font(name="Calibri", size=13, bold=True)
    inp.alignment = _left()
    inp.border = _amber_border()

    # Merge input cell across 3 columns for easier typing
    ws.merge_cells("B3:D3")

    hint = ws.cell(row=4, column=2, value="← Hier eingeben (Händler, Datum, Artikel, Quittungs-Nr. …)")
    hint.font = Font(name="Calibri", size=9, italic=True, color="6B7280")

    macro_hint = ws.cell(row=4, column=6,
                         value="Tipp: Mit dem Makro QuittungSuchen() dieses Feld direkt aufrufen (siehe vba_search.bas)")
    macro_hint.font = Font(name="Calibri", size=9, italic=True, color="4B5563")

    # ── LEFT: Quittungen ──────────────────────────────────────────────────────
    n_recv = max(len(receipts) + 5, 100)

    ws.row_dimensions[6].height = 22
    ws.merge_cells("A6:G6")
    s1 = ws["A6"]
    s1.value = "Quittungen"
    s1.font = _bold(11)
    s1.fill = _fill(COLOR_SEC_RECEIPT)
    s1.alignment = _center()
    s1.border = _border()

    recv_headers = ["Nr.", "Quittungs-Nr.", "Händler", "Datum", "Uhrzeit", "Zahlungsart", "Gesamtbetrag"]
    ws.row_dimensions[7].height = 20
    for col, h in enumerate(recv_headers, start=1):
        c = ws.cell(row=7, column=col, value=h)
        c.font = _header_font()
        c.fill = _fill(COLOR_HEADER_BG)
        c.alignment = _center()
        c.border = _border()

    # FILTER formula — searches Übersicht cols B(Quittungs-Nr), C(Händler), D(Datum), F(Zahlungsart)
    # Empty search (B3="") shows ALL rows since SEARCH("", x) = 1 always
    recv_filter = (
        f'=IFERROR(FILTER(Übersicht!A3:H{n_recv + 2},'
        f'(ISNUMBER(SEARCH(B3,Übersicht!B3:B{n_recv + 2}))'
        f'+ISNUMBER(SEARCH(B3,Übersicht!C3:C{n_recv + 2}))'
        f'+ISNUMBER(SEARCH(B3,Übersicht!D3:D{n_recv + 2}))'
        f'+ISNUMBER(SEARCH(B3,Übersicht!F3:F{n_recv + 2})))>0,'
        f'Übersicht!C3:C{n_recv + 2}<>"")'
        f',"(Keine Treffer — Suchbegriff verfeinern)")'
    )
    ws.row_dimensions[8].height = 18
    rf = ws.cell(row=8, column=1, value=recv_filter)
    rf.font = _normal(10)
    rf.alignment = _left()

    # ── RIGHT: Artikel ────────────────────────────────────────────────────────
    OFF = 9  # column I  (gap column H = col 8)
    n_art = max(sum(len(r.get("artikel") or []) for r in receipts) + 5, 100)

    ws.row_dimensions[6].height = 22
    ws.merge_cells(f"{get_column_letter(OFF)}6:{get_column_letter(OFF + 6)}6")
    s2 = ws.cell(row=6, column=OFF, value="Artikel")
    s2.font = _bold(11)
    s2.fill = _fill(COLOR_SEC_ARTICLE)
    s2.alignment = _center()
    s2.border = _border()

    art_headers = ["Quittungs-Nr.", "Händler", "Datum", "Bezeichnung", "Menge", "Einzelpreis", "Gesamtpreis"]
    ws.row_dimensions[7].height = 20
    for col, h in enumerate(art_headers, start=OFF):
        c = ws.cell(row=7, column=col, value=h)
        c.font = _header_font()
        c.fill = _fill(COLOR_HEADER_BG)
        c.alignment = _center()
        c.border = _border()

    # FILTER formula — searches Artikel-Liste cols A(Quittungs-Nr), B(Händler), C(Datum), D(Bezeichnung)
    art_filter = (
        f'=IFERROR(FILTER(\'Artikel-Liste\'!A3:G{n_art + 2},'
        f'(ISNUMBER(SEARCH(B3,\'Artikel-Liste\'!A3:A{n_art + 2}))'
        f'+ISNUMBER(SEARCH(B3,\'Artikel-Liste\'!B3:B{n_art + 2}))'
        f'+ISNUMBER(SEARCH(B3,\'Artikel-Liste\'!C3:C{n_art + 2}))'
        f'+ISNUMBER(SEARCH(B3,\'Artikel-Liste\'!D3:D{n_art + 2})))>0,'
        f'\'Artikel-Liste\'!D3:D{n_art + 2}<>"")'
        f',"(Keine Treffer — Suchbegriff verfeinern)")'
    )
    ws.row_dimensions[8].height = 18
    af = ws.cell(row=8, column=OFF, value=art_filter)
    af.font = _normal(10)
    af.alignment = _left()

    # ── Info row ──────────────────────────────────────────────────────────────
    info_row = 10 + n_recv
    ws.merge_cells(f"A{info_row}:N{info_row}")
    info = ws.cell(row=info_row, column=1,
                   value="Hinweis: Die Suchfunktion verwendet die Excel-Funktion FILTER (Excel 365 / Excel 2021+). "
                         "Für ältere Excel-Versionen bitte das VBA-Makro aus vba_search.bas importieren.")
    info.font = Font(name="Calibri", size=9, italic=True, color="6B7280")
    info.alignment = _left()

    _set_col_widths(ws, {
        "A": 5,  "B": 20, "C": 26, "D": 12, "E": 10, "F": 14, "G": 16,
        "H": 2,  # gap
        "I": 20, "J": 26, "K": 12, "L": 30, "M": 8,  "N": 13, "O": 13,
    })


# ---------------------------------------------------------------------------
# Summary Sheet  (Tab 2)
# ---------------------------------------------------------------------------
def _create_summary_sheet(wb: Workbook, receipts: list):
    ws = wb.create_sheet("Übersicht")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    ws.row_dimensions[1].height = 36
    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = f"Quittungsübersicht — erstellt am {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    title_cell.font = Font(name="Calibri", bold=True, size=14, color=COLOR_HEADER_FG)
    title_cell.fill = _fill(COLOR_HEADER_BG)
    title_cell.alignment = _center()

    headers = ["#", "Quittungs-Nr.", "Händler", "Datum", "Uhrzeit", "Zahlungsart", "Artikel (Anz.)", "Gesamtbetrag"]
    ws.row_dimensions[2].height = 22
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.font = _header_font()
        cell.fill = _fill(COLOR_HEADER_BG)
        cell.alignment = _center()
        cell.border = _border()

    grand_total = 0.0
    for idx, r in enumerate(receipts, start=1):
        row = idx + 2
        ws.row_dimensions[row].height = 20
        fill = _fill(COLOR_ROW_ALT) if idx % 2 == 0 else None

        values = [
            idx,
            r.get("quittung_nr") or "—",
            r.get("haendler") or "—",
            r.get("datum") or "—",
            r.get("uhrzeit") or "—",
            r.get("zahlungsart") or "—",
            len(r.get("artikel") or []),
            r.get("gesamtbetrag"),
        ]

        for col, val in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.font = _normal()
            cell.border = _border()
            if fill:
                cell.fill = fill
            if col == 1:
                cell.alignment = _center()
            elif col == 8:
                cell.alignment = _right()
                if val is not None:
                    cell.number_format = '#,##0.00'
                    grand_total += float(val)
            else:
                cell.alignment = _left()

    total_row = len(receipts) + 3
    ws.row_dimensions[total_row].height = 24
    ws.merge_cells(f"A{total_row}:G{total_row}")
    label = ws.cell(row=total_row, column=1, value="GESAMT")
    label.font = _bold(11)
    label.fill = _fill(COLOR_TOTAL_BG)
    label.alignment = Alignment(horizontal="right", vertical="center")
    label.border = _border()

    total_cell = ws.cell(row=total_row, column=8, value=grand_total)
    total_cell.font = _bold(12)
    total_cell.fill = _fill(COLOR_TOTAL_BG)
    total_cell.alignment = _right()
    total_cell.number_format = '#,##0.00'
    total_cell.border = _border()

    _set_col_widths(ws, {"A": 5, "B": 18, "C": 28, "D": 14, "E": 10, "F": 16, "G": 14, "H": 18})


# ---------------------------------------------------------------------------
# Artikel-Liste Sheet  (Tab 3 — consolidated article list for search)
# ---------------------------------------------------------------------------
def _create_articles_sheet(wb: Workbook, receipts: list):
    ws = wb.create_sheet("Artikel-Liste")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    ws.row_dimensions[1].height = 32
    ws.merge_cells("A1:G1")
    h = ws["A1"]
    h.value = "Alle Artikel (Suchindex)"
    h.font = Font(name="Calibri", bold=True, size=14, color=COLOR_HEADER_FG)
    h.fill = _fill(COLOR_HEADER_BG)
    h.alignment = _center()

    headers = ["Quittungs-Nr.", "Händler", "Datum", "Bezeichnung", "Menge", "Einzelpreis", "Gesamtpreis"]
    ws.row_dimensions[2].height = 22
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.font = _header_font()
        cell.fill = _fill(COLOR_HEADER_BG)
        cell.alignment = _center()
        cell.border = _border()

    row = 3
    for r in receipts:
        quittung_nr = r.get("quittung_nr") or "—"
        haendler = r.get("haendler") or "—"
        datum = r.get("datum") or "—"
        for pos, item in enumerate(r.get("artikel") or []):
            fill = _fill(COLOR_ROW_ALT) if row % 2 == 0 else None
            ws.row_dimensions[row].height = 18

            values = [
                quittung_nr,
                haendler,
                datum,
                item.get("bezeichnung") or "—",
                item.get("menge"),
                item.get("einzelpreis"),
                item.get("gesamtpreis"),
            ]

            for col, val in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=val)
                cell.font = _normal(10)
                cell.border = _border()
                if fill:
                    cell.fill = fill
                if col in (6, 7) and val is not None:
                    cell.number_format = '#,##0.00'
                    cell.alignment = _right()
                elif col == 5:
                    cell.alignment = _center()
                else:
                    cell.alignment = _left()
            row += 1

    ws.auto_filter.ref = f"A2:G{max(row - 1, 3)}"
    _set_col_widths(ws, {"A": 18, "B": 28, "C": 14, "D": 35, "E": 8, "F": 14, "G": 14})


# ---------------------------------------------------------------------------
# Individual Receipt Sheet
# ---------------------------------------------------------------------------
def _create_receipt_sheet(wb: Workbook, receipt: dict, idx: int):
    haendler = receipt.get("haendler") or f"Quittung {idx}"
    sheet_name = f"{idx}. {haendler}"[:28].rstrip()
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False

    row = 1

    ws.row_dimensions[row].height = 32
    ws.merge_cells(f"A{row}:E{row}")
    h = ws.cell(row=row, column=1, value=haendler)
    h.font = Font(name="Calibri", bold=True, size=16, color=COLOR_HEADER_FG)
    h.fill = _fill(COLOR_HEADER_BG)
    h.alignment = _center()
    row += 1

    meta_pairs = [
        ("Quittungs-Nr.", receipt.get("quittung_nr")),
        ("Kunden-Nr.", receipt.get("kunden_nr")),
        ("Datum", receipt.get("datum")),
        ("Uhrzeit", receipt.get("uhrzeit")),
        ("Zahlungsart", receipt.get("zahlungsart")),
        ("Währung", receipt.get("waehrung")),
        ("Steuersatz", receipt.get("steuersatz")),
        ("Händler Adresse", receipt.get("haendler_adresse")),
        ("Empfänger", receipt.get("empfaenger_name")),
        ("Empfänger Adresse", receipt.get("empfaenger_adresse")),
    ]
    meta_pairs = [(k, v) for k, v in meta_pairs if v is not None]

    ws.row_dimensions[row].height = 6
    row += 1

    for label, value in meta_pairs:
        ws.row_dimensions[row].height = 20
        lc = ws.cell(row=row, column=1, value=label)
        lc.font = _bold(10)
        lc.fill = _fill(COLOR_SUMMARY_BG)
        lc.alignment = _left()
        lc.border = _border()

        vc = ws.cell(row=row, column=2, value=value or "—")
        vc.font = _normal(10)
        vc.alignment = _left()
        vc.border = _border()
        row += 1

    row += 1

    col_headers = ["Pos.", "Bezeichnung", "Menge", "Einzelpreis", "Gesamtpreis"]
    ws.row_dimensions[row].height = 22
    for col, h in enumerate(col_headers, start=1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = _header_font()
        cell.fill = _fill(COLOR_HEADER_BG)
        cell.alignment = _center()
        cell.border = _border()
    row += 1

    artikel = receipt.get("artikel") or []
    for pos, item in enumerate(artikel, start=1):
        ws.row_dimensions[row].height = 18
        fill = _fill(COLOR_ROW_ALT) if pos % 2 == 0 else None

        cells_data = [
            (pos, _center()),
            (item.get("bezeichnung") or "—", _left()),
            (item.get("menge"), _center()),
            (item.get("einzelpreis"), _right()),
            (item.get("gesamtpreis"), _right()),
        ]

        for col, (val, align) in enumerate(cells_data, start=1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.font = _normal(10)
            cell.alignment = align
            cell.border = _border()
            if fill:
                cell.fill = fill
            if col in (4, 5) and val is not None:
                cell.number_format = '#,##0.00'
        row += 1

    row += 1
    totals = [
        ("Zwischensumme", receipt.get("zwischensumme")),
        (f"Steuer ({receipt.get('steuersatz') or ''})", receipt.get("steuer")),
        ("GESAMTBETRAG", receipt.get("gesamtbetrag")),
    ]

    for label, value in totals:
        ws.row_dimensions[row].height = 22
        is_grand = label.startswith("GESAMT")

        ws.merge_cells(f"A{row}:C{row}")
        lc = ws.cell(row=row, column=1, value=label)
        lc.font = _bold(11) if is_grand else _bold(10)
        lc.fill = _fill(COLOR_TOTAL_BG) if is_grand else _fill(COLOR_SUMMARY_BG)
        lc.alignment = Alignment(horizontal="right", vertical="center")
        lc.border = _border()

        ws.merge_cells(f"D{row}:E{row}")
        vc = ws.cell(row=row, column=4, value=value)
        vc.font = _bold(12) if is_grand else _normal(10)
        vc.fill = _fill(COLOR_TOTAL_BG) if is_grand else _fill(COLOR_SUMMARY_BG)
        vc.alignment = _right()
        vc.border = _border()
        if value is not None:
            vc.number_format = f'#,##0.00 "{receipt.get("waehrung") or "EUR"}"'
        row += 1

    _set_col_widths(ws, {"A": 6, "B": 35, "C": 10, "D": 16, "E": 16})
